import openpyxl as xl
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet


class ExcelManager:
    """Helpers for building xlsx workbooks from pandas DataFrames."""

    @staticmethod
    def create_xlsx() -> xl.Workbook:
        """Create a fresh, empty workbook."""
        return xl.Workbook()

    @staticmethod
    def create_worksheet(workbook: xl.Workbook, sheet_name: str) -> Worksheet:
        """Rename the workbook's default sheet and return it."""
        worksheet = workbook.active
        worksheet.title = sheet_name
        return worksheet

    @staticmethod
    def append_rows_to_worksheet(dataframe: pd.DataFrame, worksheet: Worksheet) -> None:
        """Write the DataFrame header plus every row into the worksheet."""
        worksheet.append(list(dataframe.columns))
        for row in dataframe.itertuples(index=False, name=None):
            worksheet.append(list(row))

    @staticmethod
    def assign_range_as_table(
        worksheet: Worksheet,
        table_name: str,
        num_columns: int,
        num_rows: int,
        style_name: str = "TableStyleMedium9",
    ) -> None:
        """Wrap the written range (header + rows) in a styled Excel table."""
        if num_columns < 1 or num_rows < 1:
            return

        last_col = get_column_letter(num_columns)
        last_row = num_rows + 1  # +1 for the header row
        table = Table(displayName=table_name, ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name=style_name,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    @staticmethod
    def write_dataframe_as_table(
        dataframe: pd.DataFrame,
        xlsx_path,
        sheet_name: str,
        table_name: str,
    ) -> None:
        """Create a workbook, write the DataFrame as a table, and save it."""
        workbook = ExcelManager.create_xlsx()
        worksheet = ExcelManager.create_worksheet(workbook, sheet_name)
        ExcelManager.append_rows_to_worksheet(dataframe, worksheet)
        ExcelManager.assign_range_as_table(
            worksheet, table_name, len(dataframe.columns), len(dataframe)
        )
        workbook.save(xlsx_path)
